"""Offline tests for the candidate-side screener. No network."""

from __future__ import annotations

import io

import pytest

from finetune.jobs import JOBS
from finetune.schemas import FitScore
from rehearse_core.llm.fake import FakeProvider
from screener.gap import gap_report
from screener.parse_resume import parse_resume
from screener.schemas import GapReport, ResumeProfile
from screener.score import score_fit

_RESUME = "Jane Doe. Senior engineer, 8 years. Python, Kubernetes, Terraform. BS Computer Science."


def _fit(**over: object) -> FitScore:
    base: dict[str, object] = {
        "overall_fit": 78,
        "skills_match": 4,
        "experience_match": 4,
        "seniority_match": 5,
        "matched_skills": ["Python", "Kubernetes"],
        "missing_skills": ["Go"],
        "rationale": "Solid platform background.",
    }
    base.update(over)
    return FitScore.model_validate(base)


def test_parse_resume_returns_profile() -> None:
    profile = ResumeProfile(summary="Senior platform engineer", skills=["Python", "Kubernetes"])
    provider = FakeProvider(profile)
    got = parse_resume(_RESUME, provider)
    assert got.skills == ["Python", "Kubernetes"]


def test_parse_resume_rejects_empty() -> None:
    provider = FakeProvider(ResumeProfile(summary="x", skills=["Python"]))
    with pytest.raises(ValueError):
        parse_resume("   ", provider)


def test_score_fit_uses_training_prompt_and_returns_fitscore() -> None:
    provider = FakeProvider(_fit(overall_fit=81))
    got = score_fit(_RESUME, JOBS[0], provider)
    assert got.overall_fit == 81
    # score_fit sends the same prompt the model was trained on: JD detail + resume.
    assert provider.last_user is not None and _RESUME in provider.last_user
    assert JOBS[0]["role"] in provider.last_user


def test_gap_report_flips_score_to_candidate_view() -> None:
    report = gap_report(
        _fit(matched_skills=["Python", "Kubernetes"], missing_skills=["Go", "gRPC"])
    )
    assert isinstance(report, GapReport)
    assert report.strengths == ["Python", "Kubernetes"]
    assert report.gaps == ["Go", "gRPC"]
    assert "Go" in report.summary and "gRPC" in report.summary


def test_gap_report_handles_full_match() -> None:
    report = gap_report(_fit(matched_skills=["Python"], missing_skills=[]))
    assert report.gaps == []
    assert "core skills" in report.summary


def test_extract_required_skills_returns_granular_skills() -> None:
    from screener.extract_skills import extract_required_skills
    from screener.schemas import RequiredSkills

    provider = FakeProvider(RequiredSkills(required=["FastAPI", "Railway"], nice_to_have=["Redis"]))
    got = extract_required_skills("We need FastAPI on Railway...", provider)
    assert got.required == ["FastAPI", "Railway"]


def test_extract_text_plain() -> None:
    from screener.extract_text import extract_text

    assert extract_text("resume.txt", b"Jane Doe, Python, RAG") == "Jane Doe, Python, RAG"


def test_extract_text_docx() -> None:
    import docx

    from screener.extract_text import extract_text

    doc = docx.Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Python, Kubernetes")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("resume.docx", buf.getvalue())
    assert "Jane Doe" in text and "Kubernetes" in text
